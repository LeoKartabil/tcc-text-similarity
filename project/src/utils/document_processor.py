"""
Módulo para processamento de documentos em diferentes formatos.
Suporta PDF, DOCX e TXT com diferentes encodings.
"""

import os
import logging
from typing import Optional, Union, List
from pathlib import Path
import io

# Document processing libraries
import PyPDF2
import pdfplumber
from docx import Document
import chardet

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class DocumentProcessor:
    """
    Classe responsável pelo processamento de documentos em múltiplos formatos.
    
    Suporta:
    - PDF (usando PyPDF2 e pdfplumber como fallback)
    - DOCX (usando python-docx)
    - TXT (com detecção automática de encoding)
    """
    
    SUPPORTED_FORMATS = {'.pdf', '.docx', '.txt'}
    
    def __init__(self):
        """Inicializa o processador de documentos."""
        self.processed_files = []
        
    def extract_text_from_file(self, file_path: Union[str, Path]) -> str:
        """
        Extrai texto de um arquivo baseado em sua extensão.
        
        Args:
            file_path: Caminho para o arquivo
            
        Returns:
            str: Texto extraído do arquivo
            
        Raises:
            ValueError: Se o formato não for suportado
            FileNotFoundError: Se o arquivo não existir
        """
        file_path = Path(file_path)
        
        if not file_path.exists():
            raise FileNotFoundError(f"Arquivo não encontrado: {file_path}")
            
        extension = file_path.suffix.lower()
        
        if extension not in self.SUPPORTED_FORMATS:
            raise ValueError(f"Formato não suportado: {extension}. "
                           f"Formatos suportados: {', '.join(self.SUPPORTED_FORMATS)}")
        
        try:
            if extension == '.pdf':
                return self._extract_from_pdf(file_path)
            elif extension == '.docx':
                return self._extract_from_docx(file_path)
            elif extension == '.txt':
                return self._extract_from_txt(file_path)
        except Exception as e:
            logger.error(f"Erro ao processar arquivo {file_path}: {str(e)}")
            raise
            
    def extract_text_from_uploaded_file(self, uploaded_file) -> str:
        """
        Extrai texto de um arquivo carregado via Streamlit.
        
        Args:
            uploaded_file: Arquivo carregado via st.file_uploader
            
        Returns:
            str: Texto extraído do arquivo
        """
        file_extension = Path(uploaded_file.name).suffix.lower()
        
        if file_extension not in self.SUPPORTED_FORMATS:
            raise ValueError(f"Formato não suportado: {file_extension}")
        
        try:
            if file_extension == '.pdf':
                return self._extract_from_pdf_bytes(uploaded_file.read())
            elif file_extension == '.docx':
                return self._extract_from_docx_bytes(uploaded_file.read())
            elif file_extension == '.txt':
                return self._extract_from_txt_bytes(uploaded_file.read())
        except Exception as e:
            logger.error(f"Erro ao processar arquivo carregado {uploaded_file.name}: {str(e)}")
            raise
    
    def _extract_from_pdf(self, file_path: Path) -> str:
        """
        Extrai texto de arquivo PDF usando PyPDF2 e pdfplumber como fallback.
        
        Args:
            file_path: Caminho para o arquivo PDF
            
        Returns:
            str: Texto extraído
        """
        text = ""
        
        # Primeira tentativa com PyPDF2
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                        
            if text.strip():
                logger.info(f"Texto extraído com PyPDF2 de {file_path}")
                return text.strip()
        except Exception as e:
            logger.warning(f"PyPDF2 falhou para {file_path}: {str(e)}")
        
        # Fallback com pdfplumber
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                        
            logger.info(f"Texto extraído com pdfplumber de {file_path}")
            return text.strip()
        except Exception as e:
            logger.error(f"pdfplumber também falhou para {file_path}: {str(e)}")
            raise
    
    def _extract_from_pdf_bytes(self, file_bytes: bytes) -> str:
        """
        Extrai texto de bytes de arquivo PDF.
        
        Args:
            file_bytes: Bytes do arquivo PDF
            
        Returns:
            str: Texto extraído
        """
        text = ""
        
        # Primeira tentativa com PyPDF2
        try:
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
            for page in pdf_reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
                    
            if text.strip():
                return text.strip()
        except Exception as e:
            logger.warning(f"PyPDF2 falhou para arquivo carregado: {str(e)}")
        
        # Fallback com pdfplumber
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                        
            return text.strip()
        except Exception as e:
            logger.error(f"pdfplumber também falhou para arquivo carregado: {str(e)}")
            raise
    
    def _extract_from_docx(self, file_path: Path) -> str:
        """
        Extrai texto de arquivo DOCX.
        
        Args:
            file_path: Caminho para o arquivo DOCX
            
        Returns:
            str: Texto extraído
        """
        try:
            doc = Document(file_path)
            text = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text.append(paragraph.text)
            
            # Também extrai texto de tabelas
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text.append(cell.text)
            
            result = "\n".join(text)
            logger.info(f"Texto extraído de DOCX: {file_path}")
            return result
        except Exception as e:
            logger.error(f"Erro ao extrair texto de DOCX {file_path}: {str(e)}")
            raise
    
    def _extract_from_docx_bytes(self, file_bytes: bytes) -> str:
        """
        Extrai texto de bytes de arquivo DOCX.
        
        Args:
            file_bytes: Bytes do arquivo DOCX
            
        Returns:
            str: Texto extraído
        """
        try:
            doc = Document(io.BytesIO(file_bytes))
            text = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text.append(paragraph.text)
            
            # Também extrai texto de tabelas
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text.append(cell.text)
            
            result = "\n".join(text)
            logger.info("Texto extraído de DOCX carregado")
            return result
        except Exception as e:
            logger.error(f"Erro ao extrair texto de DOCX carregado: {str(e)}")
            raise
    
    def _extract_from_txt(self, file_path: Path) -> str:
        """
        Extrai texto de arquivo TXT com detecção automática de encoding.
        
        Args:
            file_path: Caminho para o arquivo TXT
            
        Returns:
            str: Texto extraído
        """
        # Detecta encoding
        with open(file_path, 'rb') as file:
            raw_data = file.read()
            encoding_result = chardet.detect(raw_data)
            encoding = encoding_result['encoding'] or 'utf-8'
        
        # Lê arquivo com encoding detectado
        try:
            with open(file_path, 'r', encoding=encoding) as file:
                text = file.read()
            logger.info(f"Texto extraído de TXT: {file_path} (encoding: {encoding})")
            return text
        except UnicodeDecodeError:
            # Fallback para utf-8 com ignore
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as file:
                text = file.read()
            logger.warning(f"Usado fallback UTF-8 para {file_path}")
            return text
    
    def _extract_from_txt_bytes(self, file_bytes: bytes) -> str:
        """
        Extrai texto de bytes de arquivo TXT.
        
        Args:
            file_bytes: Bytes do arquivo TXT
            
        Returns:
            str: Texto extraído
        """
        # Detecta encoding
        encoding_result = chardet.detect(file_bytes)
        encoding = encoding_result['encoding'] or 'utf-8'
        
        try:
            text = file_bytes.decode(encoding)
            logger.info(f"Texto extraído de TXT carregado (encoding: {encoding})")
            return text
        except UnicodeDecodeError:
            # Fallback para utf-8 com ignore
            text = file_bytes.decode('utf-8', errors='ignore')
            logger.warning("Usado fallback UTF-8 para arquivo TXT carregado")
            return text
    
    def get_file_info(self, file_path: Union[str, Path]) -> dict:
        """
        Obtém informações sobre um arquivo.
        
        Args:
            file_path: Caminho para o arquivo
            
        Returns:
            dict: Informações do arquivo
        """
        file_path = Path(file_path)
        
        if not file_path.exists():
            raise FileNotFoundError(f"Arquivo não encontrado: {file_path}")
        
        stat = file_path.stat()
        
        return {
            'name': file_path.name,
            'size': stat.st_size,
            'extension': file_path.suffix.lower(),
            'is_supported': file_path.suffix.lower() in self.SUPPORTED_FORMATS,
            'modified_time': stat.st_mtime
        }
    
    def validate_file(self, file_path: Union[str, Path]) -> bool:
        """
        Valida se um arquivo pode ser processado.
        
        Args:
            file_path: Caminho para o arquivo
            
        Returns:
            bool: True se o arquivo pode ser processado
        """
        try:
            file_info = self.get_file_info(file_path)
            return file_info['is_supported']
        except Exception:
            return False
    
    def batch_process(self, file_paths: List[Union[str, Path]]) -> dict:
        """
        Processa múltiplos arquivos em lote.
        
        Args:
            file_paths: Lista de caminhos para arquivos
            
        Returns:
            dict: Resultados do processamento
        """
        results = {
            'successful': {},
            'failed': {},
            'total_processed': 0
        }
        
        for file_path in file_paths:
            try:
                text = self.extract_text_from_file(file_path)
                results['successful'][str(file_path)] = text
                results['total_processed'] += 1
                logger.info(f"Processado com sucesso: {file_path}")
            except Exception as e:
                results['failed'][str(file_path)] = str(e)
                logger.error(f"Falha ao processar: {file_path} - {str(e)}")
        
        return results